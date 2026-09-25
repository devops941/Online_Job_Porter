"""Excel and PDF report generation (openpyxl + ReportLab)."""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def to_excel(title: str, columns: list[str], rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31] or "Report"

    sheet.append(columns)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(list(row))

    for idx, column in enumerate(columns, start=1):
        width = max(len(str(column)), *(len(str(r[idx - 1])) for r in rows)) if rows else len(column)
        sheet.column_dimensions[get_column_letter(idx)].width = min(max(width + 4, 12), 60)

    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(title: str, columns: list[str], rows: list[list], summary: dict | None = None) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    heading = ParagraphStyle(
        "Heading", parent=styles["Heading1"], fontSize=18, spaceAfter=8, textColor=colors.HexColor("#111827")
    )
    small = ParagraphStyle("Small", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#6B7280"))

    story = [Paragraph(title, heading), Paragraph(f"Generated {datetime.now():%Y-%m-%d %H:%M}", small)]
    if summary:
        story += [Spacer(1, 8), Paragraph(" · ".join(f"{k}: {v}" for k, v in summary.items()), small)]
    story.append(Spacer(1, 14))

    data = [columns] + [[str(c) if c is not None else "" for c in r] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buffer.getvalue()
